"""Seeded document fuzzing plus independent upstream fixtures. No network at runtime.
External fixtures: python-openxml/python-docx tests/test_files (download separately).
Usage: python tests/randomized.py [--qa]
"""
import hashlib
import json
import os
import pathlib
import random
import subprocess
import sys
import zipfile
import fitz
from docx import Document
from docx.shared import Cm, Pt
ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "randomized"
OUT.mkdir(parents=True, exist_ok=True)
RNG = random.Random(20261001)
results = []


def cli(tool, *args, expect=0):
    proc = subprocess.run([sys.executable, str(ROOT/"tools"/tool), *map(str,args)], cwd=ROOT,
        capture_output=True, text=True, encoding="utf-8", timeout=360,
        env={**os.environ,"PYTHONUTF8":"1"})
    if proc.returncode != expect:
        raise AssertionError(f"{tool}: exit {proc.returncode}\n{proc.stdout}\n{proc.stderr}")
    return proc.stdout + proc.stderr


def edit(src, dst, ops):
    digest = hashlib.sha256(src.read_bytes()).hexdigest()
    plan = dst.with_suffix(".json")
    plan.write_text(json.dumps(ops,ensure_ascii=False),encoding="utf-8")
    cli("docx_edit.py",src,dst,"--ops",plan)
    assert hashlib.sha256(src.read_bytes()).hexdigest() == digest
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(dst) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "word/document.xml": assert a.read(name) == b.read(name), name
    return Document(dst)


image = OUT/"chart.png"
pix = fitz.Pixmap(fitz.csRGB,fitz.IRect(0,0,200,100),False); pix.clear_with(160); pix.save(image)
qa_pairs = []
for i in range(20):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(RNG.choice([2,2.2,2.5])); sec.right_margin = Cm(2)
    style = doc.styles['Normal']; style.font.name=RNG.choice(['Arial','Calibri','Times New Roman']); style.font.size=Pt(RNG.choice([10,11,12]))
    doc.add_heading(f"Случайный отчёт {i}",1)
    p=doc.add_paragraph(); p.add_run("Срок "); p.add_run("оплаты").bold=True; p.add_run(" — 30 дней.")
    doc.add_paragraph("Опорный абзац для вставки.")
    doc.add_paragraph("Перемещаемый абзац.")
    for j in range(RNG.randrange(2,8)):
        doc.add_paragraph(f"Контекст {j}. " + "Данные проверяются после изменения документа. "*RNG.randrange(1,5))
    doc.add_picture(str(image),width=Cm(RNG.choice([4,6,8])))
    cols, rows = RNG.randrange(2,6), RNG.randrange(3,12)
    table=doc.add_table(rows=rows,cols=cols); table.style=RNG.choice(['Table Grid','Light Shading Accent 1','Light List Accent 1'])
    for ri,row in enumerate(table.rows):
        for ci,cell in enumerate(row.cells):
            cell.text=f"Колонка {ci}" if ri==0 else f"{ri}:{ci}"
    doc.add_heading("Удаляемый раздел",1); doc.add_paragraph("Удаляемое содержимое.")
    doc.add_heading("Сохраняемый раздел",1); doc.add_paragraph("Финальный контрольный текст.")
    doc.sections[0].header.paragraphs[0].text=f"Колонтитул {i}"
    src,dst=OUT/f"seed-{i:02}.docx",OUT/f"seed-{i:02}.edited.docx"
    doc.save(src)
    nadd=RNG.randrange(1,6)
    values=[f"Добавлено {i}"] + [str(RNG.randrange(1,999)) for _ in range(cols-1)]
    ops=[{"op":"replace","find":"оплаты — 30 дней","replace":"оплаты — 45 дней"},
         {"op":"cell","table":0,"row":1,"col":1,"text":"Изменено"},
         {"op":"set_text","block":"Опорный абзац","text":"Новая опора."},
         {"op":"insert","after":"Новая опора","texts":["Новый первый.","Новый второй."]},
         {"op":"move","block":"Перемещаемый абзац","after":"Новый первый"},
         {"op":"insert_table","after":"Новый второй","rows":[["Ключ","Значение"],["A","12"]]},
         {"op":"delete_section","heading":"Удаляемый раздел"},
         {"op":"delete","block":"Новый второй"},
         {"op":"resize_image","image":0,"width_cm":3},
         {"op":"page_break_before","block":"Сохраняемый раздел"}]
    # New table inserted before the original: original table now has index 1.
    ops += [{"op":"add_row","table":1,"after":-1,"values":values} for _ in range(nadd)]
    ops += [{"op":"del_row","table":1,"row":2}]
    edited=edit(src,dst,ops)
    text='\n'.join(p.text for p in edited.paragraphs)
    assert "45 дней" in text and "30 дней" not in text
    assert "Удаляемый раздел" not in text and "Удаляемое содержимое" not in text
    assert "Новый второй" not in text and "Финальный контрольный текст" in text
    order=[p.text for p in edited.paragraphs]
    assert order.index("Перемещаемый абзац.")==order.index("Новый первый.")+1
    assert len(edited.tables)==2 and len(edited.tables[1].rows)==rows+nadd-1
    assert edited.tables[1].cell(1,1).text=="Изменено"
    assert edited.tables[1].rows[-1].cells[0].text==values[0]
    assert abs(edited.inline_shapes[0].width.cm-3)<.001 and abs(edited.inline_shapes[0].height.cm-1.5)<.001
    assert edited.sections[0].header.paragraphs[0].text==f"Колонтитул {i}"
    assert any(r.bold and r.text=='оплаты' for p in edited.paragraphs for r in p.runs)
    results.append({"file":dst.name,"rows":rows,"columns":cols,"added":nadd,"status":"PASS"})
    print(f"PASS random {i}: {rows}x{cols}, +{nadd} rows",flush=True)
    if i in (2,11,19): qa_pairs.append((src,dst))

external=OUT/"external"
plans={
 "test.docx":[{"op":"replace","find":"python-docx was here!","replace":"python-docx was there!"}],
 "blk-inner-content.docx":[{"op":"replace","find":"P1","replace":"Intro"},{"op":"cell","table":0,"row":0,"col":1,"text":"Edited"},{"op":"add_row","table":0,"after":1,"values":["Added","42"]}],
 "sct-inner-content.docx":[{"op":"cell","table":1,"row":0,"col":0,"text":"Edited section table"},{"op":"insert","after":"P8","text":"Inserted paragraph."}],
 "having-images.docx":[{"op":"resize_image","image":i,"width_cm":round(s.width.cm*.65,4)} for i,s in enumerate(Document(external/"having-images.docx").inline_shapes)],
}
for name,ops in plans.items():
    src,dst=external/name,external/(pathlib.Path(name).stem+".edited.docx")
    edited=edit(src,dst,ops)
    if name=='having-images.docx':
        original=Document(src)
        assert len(edited.inline_shapes)==len(original.inline_shapes)==5
        for a,b in zip(original.inline_shapes,edited.inline_shapes):
            assert abs(b.width.cm/a.width.cm-.65)<.001
            assert abs(b.width/b.height-a.width/a.height)<.001
    elif name=='test.docx': assert edited.paragraphs[0].text=='python-docx was there!'
    elif name=='blk-inner-content.docx': assert len(edited.tables[0].rows)==3 and edited.tables[0].cell(0,1).text=='Edited'
    else: assert edited.tables[1].cell(0,0).text=='Edited section table'
    results.append({"file":dst.name,"source":"python-openxml/python-docx","status":"PASS"})
    qa_pairs.append((src,dst)); print(f"PASS external {name}",flush=True)

paper=ROOT/"out/validation/attention.pdf"
pdfout=external/"attention.edited.pdf"
digest=hashlib.sha256(paper.read_bytes()).hexdigest()
cli("pdf_edit.py",paper,pdfout,"--replace","Attention Is All You Need","Attention Is All We Need")
assert hashlib.sha256(paper.read_bytes()).hexdigest()==digest
with fitz.open(paper) as before,fitz.open(pdfout) as after:
    assert len(before)==len(after)==15
    assert "Attention Is All We Need" in after[0].get_text()
    assert "Attention Is All You Need" not in after[0].get_text()
    for i in range(1,len(before)):
        assert before[i].get_text()==after[i].get_text()
        assert before[i].get_pixmap(matrix=fitz.Matrix(.5,.5)).samples==after[i].get_pixmap(matrix=fitz.Matrix(.5,.5)).samples
results.append({"file":pdfout.name,"pages":15,"status":"PASS"})
print("PASS external PDF: title changed, other 14 pages pixel-identical",flush=True)
qa_pairs.append((paper,pdfout))
if '--qa' in sys.argv:
    for src,dst in qa_pairs:
        print(f"QA {dst.name}",flush=True)
        qa=cli("qa.py",dst,"--original",src,"--allow-reflow")
        (dst.with_suffix('.qa.log')).write_text(qa,encoding='utf-8')
        print(qa.splitlines()[-1],flush=True)
(OUT/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(f"ALL GREEN: {len(results)} independent/randomized document edits",flush=True)
