"""Deterministic structural and transactional regressions (no network)."""
import pathlib
import sys
import tempfile
import unittest
import zipfile
import subprocess
import json
from unittest.mock import patch
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "tools"))
import pymupdf as fitz
from docx import Document
from docx.shared import Cm
from lxml import etree
from docx_edit import Editor, OpError, add_row, q, run_ops, set_para_text, resize_image
from safe_output import staged_output
from pdf_objects import apply_ops
from docx_kit import new_document, add_table
ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "regression"
OUT.mkdir(parents=True, exist_ok=True)


class Regression(unittest.TestCase):
    def test_atomic_rollback_and_alias(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = pathlib.Path(d)/"in", pathlib.Path(d)/"out"
            src.write_bytes(b"original"); dst.write_bytes(b"previous")
            with self.assertRaises(RuntimeError):
                with staged_output(src, dst) as temp:
                    temp.write_bytes(b"partial")
                    raise RuntimeError("validation failure")
            self.assertEqual(dst.read_bytes(), b"previous")
            self.assertEqual(src.read_bytes(), b"original")
            self.assertEqual(len(list(pathlib.Path(d).iterdir())), 2)
            with self.assertRaises(ValueError):
                with staged_output(src, src): pass

    def test_row_does_not_copy_vertical_merge(self):
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        table.cell(0,0).merge(table.cell(1,0))
        row = add_row(table._tbl, 1, ["new", "value"], Editor())
        self.assertFalse(list(row.iter(q("vMerge"))))
        self.assertEqual(len(table.rows), 3)

    def test_mixed_run_keeps_drawing(self):
        p = etree.fromstring(f'<w:p xmlns:w="{q("p").split("}")[0][1:]}"><w:r><w:t>caption</w:t><w:drawing/></w:r></w:p>')
        set_para_text(p, "new", Editor())
        self.assertEqual(len(list(p.iter(q("drawing")))), 1)
        self.assertEqual("".join(p.itertext()), "new")

    def test_bad_batch_keeps_existing_result(self):
        doc = Document(); doc.add_paragraph("unique anchor")
        src, dst, ops = OUT/"input.docx", OUT/"unchanged.docx", OUT/"bad.json"
        doc.save(src); dst.write_bytes(b"previous result")
        ops.write_text(json.dumps([{"op":"set_text", "block":"unique anchor", "text":"changed"}, {"op":"insert_table", "after":"changed", "rows":[]}]), encoding="utf-8")
        r = subprocess.run([sys.executable, str(ROOT/"tools/docx_edit.py"), str(src), str(dst), "--ops", str(ops)], capture_output=True)
        self.assertNotEqual(r.returncode, 0)
        self.assertEqual(dst.read_bytes(), b"previous result")
        self.assertEqual(Document(src).paragraphs[0].text, "unique anchor")

    def test_drawing_resize_and_render_fixture(self):
        image = OUT/"image.png"
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0,0,160,80), False)
        pix.clear_with(190); pix.save(image)
        doc = new_document(); doc.add_heading("Проверка объектов и таблиц", 1)
        doc.add_paragraph("Документ создан с нуля. Изображение ниже уменьшено, таблица расширена.", style="Body Text")
        doc.add_picture(str(image), width=Cm(8))
        add_table(doc, [["Показатель", "Значение"], ["Выручка", "120"]])
        src, dst = OUT/"objects.docx", OUT/"objects.edited.docx"
        doc.save(src)
        ops = OUT/"objects.json"
        ops.write_text(json.dumps([{"op":"resize_image", "image":0,"width_cm":6}, {"op":"add_row","table":0,"after":1,"values":["Прибыль","35"]}]), encoding="utf-8")
        r = subprocess.run([sys.executable, str(ROOT/"tools/docx_edit.py"),str(src),str(dst),"--ops",str(ops)], capture_output=True)
        self.assertEqual(r.returncode,0,r.stderr.decode(errors="replace"))
        edited = Document(dst)
        self.assertAlmostEqual(edited.inline_shapes[0].width.cm,6, places=3)
        self.assertAlmostEqual(edited.inline_shapes[0].height.cm,3, places=3)
        self.assertEqual(len(edited.tables[0].rows),3)
        with zipfile.ZipFile(src) as a, zipfile.ZipFile(dst) as b:
            for name in a.namelist():
                if name != "word/document.xml": self.assertEqual(a.read(name),b.read(name),name)
        tree = etree.fromstring(zipfile.ZipFile(src).read("word/document.xml"))
        with self.assertRaises(OpError): resize_image(tree, {"image":0,"width_cm":100},Editor())
        with self.assertRaises(OpError): resize_image(tree, {"image":0,"width_cm":6},Editor("Author"))

    def test_pdf_objects_unicode_and_overlap(self):
        src, dst = OUT/"objects.pdf", OUT/"objects.edited.pdf"
        doc = fitz.open(); doc.new_page(width=595,height=842)
        doc.save(src)
        apply_ops(doc,[{"op":"insert_text","rect":[50,50,540,100],"text":"Проверка объектов PDF"}, {"op":"insert_table","rect":[50,140,540,350],"rows":[["Показатель","Значение"],["Выручка","120"],["Прибыль","35"]]}],OUT)
        with self.assertRaises(ValueError): apply_ops(doc,[{"op":"insert_text","rect":[50,50,540,100],"text":"overlap"}],OUT)
        doc.save(dst)
        with fitz.open(dst) as check:
            self.assertIn("Проверка", check[0].get_text())
            self.assertIn("Прибыль", check[0].get_text())
        apply_ops(doc,[{"op":"duplicate_page"},{"op":"rotate_page","page":2,"angle":90},{"op":"delete_page","page":1}],OUT)
        self.assertEqual(len(doc),1); self.assertEqual(doc[0].rotation,90)
        with self.assertRaises(ValueError): apply_ops(doc,[{"op":"delete_page"}],OUT)
        doc.close()

    def test_pdf_replace_refuses_mixed_style(self):
        src, dst = OUT/"mixed.pdf", OUT/"mixed.edited.pdf"
        dst.unlink(missing_ok=True)
        doc = fitz.open(); page = doc.new_page()
        page.insert_text((50,80), "Alpha", fontsize=12)
        page.insert_text((81,80), "Beta", fontsize=12, color=(1,0,0))
        doc.save(src); doc.close()
        r = subprocess.run([sys.executable, str(ROOT/"tools/pdf_edit.py"), str(src), str(dst), "--replace", "AlphaBeta", "Other"], capture_output=True)
        self.assertNotEqual(r.returncode, 0)
        self.assertFalse(dst.exists())

    def test_pdf_replace_keeps_unrelated_text(self):
        """Redaction must not take neighbouring runs with it: a page number in a second face stays."""
        src, dst = OUT/"collateral.pdf", OUT/"collateral.edited.pdf"
        dst.unlink(missing_ok=True)
        doc = fitz.open(); page = doc.new_page()
        page.insert_text((50, 80), "ReplaceMe", fontsize=12)
        page.insert_text((50, 780), "Page 1 of 1", fontsize=9, fontname="hebo")  # bold face, no Cyrillic
        doc.save(src); doc.close()
        r = subprocess.run([sys.executable, str(ROOT/"tools/pdf_edit.py"), str(src), str(dst),
                            "--replace", "ReplaceMe", "Changed"], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        with fitz.open(dst) as d:
            text = d[0].get_text()
        self.assertIn("Changed", text)
        self.assertIn("Page 1 of 1", text)

    def test_pdf_replace_uses_system_font_when_subset_lacks_glyphs(self):
        """A subset without Cyrillic is not a dead end: the same family from the system is embedded."""
        src, dst = OUT/"subset.pdf", OUT/"subset.edited.pdf"
        dst.unlink(missing_ok=True)
        doc = fitz.open(); page = doc.new_page()
        page.insert_text((50, 80), "Ivanov Ivanovich", fontsize=12)  # embeds a Latin-only subset
        doc.save(src, garbage=4, deflate=True); doc.close()
        r = subprocess.run([sys.executable, str(ROOT/"tools/pdf_edit.py"), str(src), str(dst),
                            "--replace", "Ivanov Ivanovich", "Иванов Иван"], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("system:", r.stdout)  # the subset could not be used, the system family was
        with fitz.open(dst) as d:
            self.assertIn("Иванов Иван", d[0].get_text())

    def test_pdf_replace_grows_right_aligned_text_leftwards(self):
        """A wider number may grow into the free space left of a right-aligned column edge."""
        src, dst = OUT/"grow.pdf", OUT/"grow.edited.pdf"
        dst.unlink(missing_ok=True)
        doc = fitz.open(); page = doc.new_page()
        page.insert_text((50, 80), "Balance", fontsize=10)
        page.insert_text((300, 80), "300,85", fontsize=10)
        page.insert_text((50, 100), "Balance", fontsize=10)
        page.insert_text((300, 100), "300,85", fontsize=10)
        doc.save(src); doc.close()
        r = subprocess.run([sys.executable, str(ROOT/"tools/pdf_edit.py"), str(src), str(dst),
                            "--replace", "300,85", "500 000,00", "2"], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        with fitz.open(dst) as d:
            right = [w[2] for w in d[0].get_text("words") if w[4] == "500"]
        self.assertEqual(len(right), 2)
        self.assertAlmostEqual(right[0], right[1], delta=1.0)  # still flush right

    def test_stale_render_is_not_used(self):
        import render
        with tempfile.TemporaryDirectory() as d:
            folder = pathlib.Path(d); src = folder/"input.docx"
            Document().save(src)
            target = folder/"render"; target.mkdir()
            stale = target/"input.pdf"; stale.write_bytes(b"stale PDF")
            with patch.object(render, "to_pdf", side_effect=ValueError("conversion failed")):
                with self.assertRaises(ValueError): render.render(src, target)
            self.assertEqual(stale.read_bytes(), b"stale PDF")

    def test_tracked_move_into_itself_refused(self):
        doc = Document(); doc.add_paragraph("self anchor")
        tree = etree.fromstring(doc._element.xml.encode())
        with self.assertRaises(OpError):
            run_ops([{"op":"move", "block":"self anchor", "after":"self anchor"}], {"word/document.xml":tree}, Editor("Reviewer"), {}, set())
        self.assertFalse(list(tree.iter(q("del"))))
        body = tree.find(q("body"))
        p = next(body.iter(q("p")))
        etree.SubElement(next(p.iter(q("r"))), q("drawing"))
        target = etree.Element(q("p")); r = etree.SubElement(target,q("r")); etree.SubElement(r,q("t")).text="destination"
        body.insert(1,target)
        with self.assertRaisesRegex(OpError,"tracked move of drawings"):
            run_ops([{"op":"move","block":"#0","after":"#1"}], {"word/document.xml":tree}, Editor("Reviewer"), {}, set())
        self.assertFalse(list(tree.iter(q("del"))))

    def test_contact_sheet_respects_page_rotation(self):
        from render import sheet
        doc=fitz.open(); page=doc.new_page(width=300,height=400)
        page.draw_rect(fitz.Rect(50,60,150,80),color=None,fill=(0,0,0))
        page.set_rotation(90)
        path=OUT/'rotated-sheet.png'; sheet(doc,path,cols=1)
        pix=fitz.Pixmap(path)
        dark=[]
        for y in range(pix.height):
            for x in range(pix.width):
                start=(y*pix.width+x)*pix.n
                if max(pix.samples[start:start+3])<30: dark.append((x,y))
        self.assertTrue(dark)
        width=max(x for x,y in dark)-min(x for x,y in dark)
        height=max(y for x,y in dark)-min(y for x,y in dark)
        self.assertGreater(height,width*3,'rotated marker must be vertical in the contact sheet')
        self.assertGreater(sum(x for x,y in dark)/len(dark), pix.width/2, 'clockwise rotation must put marker on the right')
        self.assertLess(sum(y for x,y in dark)/len(dark), pix.height/2, 'clockwise rotation must put marker near the top')
        doc.close()

    def test_pdf_overflow_refused(self):
        doc = fitz.open(); doc.new_page(width=200,height=200)
        with self.assertRaises(ValueError):
            apply_ops(doc,[{"op":"insert_table","rect":[10,10,30,20],"rows":[["Too much content","Another column"],["body","value"]]}],OUT)
        doc.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
