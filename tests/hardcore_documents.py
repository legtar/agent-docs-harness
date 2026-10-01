"""Large realistic DOCX scenarios and verification of the resulting XML/package."""
import hashlib, json, pathlib, subprocess, sys, zipfile
from lxml import etree
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.shared import Cm
ROOT=pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tools'))
from docx_kit import new_document,add_table,add_toc,polish_table
from docx_edit import W,q
OUT=ROOT/'out/hardcore/documents';OUT.mkdir(parents=True,exist_ok=True)


def execute(src,dst,ops,track=None):
    before=hashlib.sha256(src.read_bytes()).hexdigest()
    plan=dst.with_suffix('.json');plan.write_text(json.dumps(ops,ensure_ascii=False,indent=2),encoding='utf-8')
    args=[sys.executable,str(ROOT/'tools/docx_edit.py'),str(src),str(dst),'--ops',str(plan)]
    if track:args+=['--track',track]
    proc=subprocess.run(args,capture_output=True,text=True,encoding='utf-8',cwd=ROOT,timeout=240)
    if proc.returncode:raise AssertionError(proc.stdout+proc.stderr)
    assert hashlib.sha256(src.read_bytes()).hexdigest()==before
    with zipfile.ZipFile(src) as a,zipfile.ZipFile(dst) as b:
        assert a.namelist()==b.namelist()
        for name in a.namelist():
            if name!='word/document.xml':assert a.read(name)==b.read(name),name
    print(f'PASS {dst.name}: {len(ops)} operations, original intact, unrelated parts identical',flush=True)


if '--complex-only' not in sys.argv:
    # Report with TOC, 180 contracts, 24 sections, and a 128-operation editing batch.
    doc=new_document();doc.add_heading('Сводный отчёт: договоры и контроль качества',0);add_toc(doc)
    doc.add_paragraph('Изначально обработано 180 договоров.',style='Body Text')
    p=doc.add_paragraph(style='Body Text');p.add_run('Срок ');p.add_run('оплаты').bold=True;p.add_run(' составляет 30 дней.')
    caption=doc.add_paragraph('Таблица 1. Реестр договоров',style='Table Caption');caption.paragraph_format.page_break_before=True
    rows=[['Идентификатор','Описание договора','Стоимость, тыс. ₽','Статус']]
    rows += [[str(i),f'Договор № {i}: поставка оборудования и техническое обслуживание регионального филиала.',str(100+i*3),'Исполнен'] for i in range(1,181)]
    add_table(doc,rows,widths=[1,6,2,2])
    body='Показатели сопоставлены с первичными документами. Контрольные значения подтверждены ответственным подразделением. Выявленные отклонения отражены в реестре и направлены на повторную проверку.'
    for i in range(1,25):
        p=doc.add_heading(f'Раздел {i}. Контроль направления',1);p.paragraph_format.page_break_before=True
        for j in range(4):doc.add_paragraph(f'Этап {j+1}. '+body,style='Body Text')
        doc.add_heading(f'Детали направления {i}',2);doc.add_paragraph('Результаты проверки зафиксированы в журнале качества.',style='Body Text')
    doc.add_paragraph('Итоговая контрольная запись.',style='Body Text')
    src,dst=OUT/'large-report.docx',OUT/'large-report.edited.docx';doc.save(src)
    ops=[{'op':'replace','find':'180 договоров','replace':'280 договоров'},
         {'op':'replace','find':'оплаты составляет 30 дней','replace':'оплаты составляет 45 дней'},
         {'op':'cell','table':'Идентификатор','row':12,'col':2,'text':'999'},
         {'op':'insert','after':'Срок оплаты','text':'Дополнительная методика.'},
         {'op':'insert_table','after':'Дополнительная методика.','rows':[['Метрика','Значение'],['Проверено','280']]},
         {'op':'set_text','block':'Итоговая контрольная запись.','text':'Итоговая контрольная запись: проверено.'},
         {'op':'delete_section','heading':'Раздел 10. Контроль направления'},
         {'op':'move','block':'Дополнительная методика.','after':'Раздел 8. Контроль направления'}]
    ops += [{'op':'add_row','table':'Идентификатор','after':-1,'values':[str(i),f'Новый договор № {i}: обслуживание филиала.',str(100+i),'Добавлен']} for i in range(181,291)]
    ops += [{'op':'del_row','table':'Идентификатор','row':1} for _ in range(10)]
    execute(src,dst,ops)
    changed=Document(dst);registry=next(t for t in changed.tables if 'Идентификатор' in t.cell(0,0).text)
    assert len(registry.rows)==281 and registry.rows[-1].cells[0].text=='290'
    assert not any(p.text.startswith('Раздел 10.') for p in changed.paragraphs)
    assert any(r.bold and r.text=='оплаты' for p in changed.paragraphs for r in p.runs)
    assert any('45 дней' in p.text for p in changed.paragraphs)
    tracked=OUT/'large-report.tracked.docx'
    execute(src,tracked,[{'op':'replace','find':'30 дней','replace':'60 дней'}, {'op':'insert','after':'Срок оплаты','text':'Вставлено с ревизией.'}, {'op':'cell','table':0,'row':8,'col':2,'text':'777'}, {'op':'add_row','table':0,'after':-1,'values':['181','Дополнительный договор','999','Ревизия']}],'Hardcore auditor')
    with zipfile.ZipFile(tracked) as z:
        root=etree.fromstring(z.read('word/document.xml'))
        assert list(root.iter(q('ins'))) and list(root.iter(q('del')))
        assert not root.xpath('//w:ins//w:ins | //w:ins//w:del',namespaces={'w':W})
    
    
# Merged cells, image in second cell paragraph, nested table, textbox, content control, sections.
image=ROOT/'out/hardcore/marker.png'
doc=new_document();doc.add_heading('Объекты и сложные таблицы',1)
t=add_table(doc,[['Группа','Показатель','Значение'],['A','Выручка','100'],['','Прибыль','30'],['','Маржа','30 %'],['B','Выручка','80']])
t.cell(1,0).merge(t.cell(3,0)).text='Группа A'
t.cell(0,0).merge(t.cell(0,1)).text='Показатели группы';polish_table(t,9300,widths=[2,5,2])
doc.add_paragraph('Карточка документа',style='Heading 2')
card=add_table(doc,[['Поле','Содержимое'],['Инструкция','Вложенные показатели']],widths=[1,2])
cell=card.cell(1,0);cell.add_paragraph().add_run().add_picture(str(image),width=Cm(1.2))
inner=card.cell(1,1).add_table(rows=2,cols=2);inner.cell(0,0).text='Вложенный ключ';inner.cell(0,1).text='Число';inner.cell(1,0).text='Лимит';inner.cell(1,1).text='17';polish_table(inner,3500,widths=[2,1])
doc.add_paragraph('Якорь портретного раздела.',style='Body Text')
# A valid inline VML textbox: replacement must preserve the shape while changing nested text.
p=doc.add_paragraph();pict=OxmlElement('w:pict');p.add_run()._r.append(pict)
V='urn:schemas-microsoft-com:vml'
shape=etree.SubElement(pict,f'{{{V}}}rect',id='hardcore_textbox',style='width:240pt;height:34pt')
box=etree.SubElement(shape,f'{{{V}}}textbox');content=etree.SubElement(box,q('txbxContent'))
bp=etree.SubElement(content,q('p'));r=etree.SubElement(bp,q('r'));etree.SubElement(r,q('t')).text='Текст внутри фигуры: исходный.'
sdt=OxmlElement('w:sdt');sdtcontent=OxmlElement('w:sdtContent');sdt.append(sdtcontent)
sp=OxmlElement('w:p');sr=OxmlElement('w:r');st=OxmlElement('w:t');st.text='Контент-контрол: исходный.';sr.append(st);sp.append(sr);sdtcontent.append(sp)
doc._element.body.insert(len(doc._element.body)-1,sdt)
sec=doc.add_section();sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Cm(29.7);sec.page_height=Cm(21)
doc.add_heading('Альбомный раздел',1);doc.add_paragraph('Широкая матрица показателей.',style='Body Text')
add_table(doc,[['Показатель','I кв.','II кв.','III кв.','IV кв.','Итого'],['Выручка','10','20','30','40','100']],widths=[3,1,1,1,1,1])
src,dst=OUT/'complex-objects.docx',OUT/'complex-objects.edited.docx';doc.save(src)
ops=[{'op':'cell','table':0,'row':3,'col':0,'text':'Группа A — проверено'},
     {'op':'del_row','table':0,'row':1},
     {'op':'add_row','table':0,'after':-1,'values':['C','Выручка','95']},
     {'op':'cell','table':1,'row':1,'col':0,'text':'Инструкция обновлена'},
     {'op':'cell','table':2,'row':1,'col':1,'text':'25'},
     {'op':'resize_image','image':0,'width_cm':1.5},
     {'op':'replace','find':'Текст внутри фигуры: исходный.','replace':'Текст внутри фигуры: обновлён.'},
     {'op':'replace','find':'Контент-контрол: исходный.','replace':'Контент-контрол: обновлён.'},
     {'op':'insert_table','after':'Якорь портретного раздела.','rows':[['Ключ','Значение'],['Проверка','Пройдена']]}]
execute(src,dst,ops)
changed=Document(dst)
assert changed.tables[0].cell(1,0).text=='Группа A — проверено'
assert len(changed.inline_shapes)==1 and abs(changed.inline_shapes[0].width.cm-1.5)<.001
with zipfile.ZipFile(src) as a,zipfile.ZipFile(dst) as b:
    before,after=etree.fromstring(a.read('word/document.xml')),etree.fromstring(b.read('word/document.xml'))
    old_sections=[etree.tostring(s) for s in before.iter(q('sectPr'))]
    new_sections=[etree.tostring(s) for s in after.iter(q('sectPr'))]
    assert old_sections==new_sections
    assert len(list(after.iter(q('pict'))))==1
    assert 'Текст внутри фигуры: обновлён.' in ''.join(after.itertext())
    assert 'Контент-контрол: обновлён.' in ''.join(after.itertext())
    assert '25' in ''.join(after.itertext())
print('ALL GREEN: complex object/section document' if '--complex-only' in sys.argv else 'ALL GREEN: large report, tracked report and complex object/section document',flush=True)
