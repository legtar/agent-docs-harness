"""Deterministic tests of the editing features (no network): DOCX columns, merges, formatting,
pictures, nested addressing, odd packages; in-place PDF rows, cells, blocks, vertical space.

    python tests/features.py
"""
import io
import json
import os
import pathlib
import subprocess
import sys
import unittest
import zipfile

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import pikepdf  # noqa: E402
import pymupdf as fitz  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from lxml import etree  # noqa: E402

import docx_edit as de  # noqa: E402
import pdf_flow as pf  # noqa: E402
from docx_kit import add_table, new_document  # noqa: E402
from render import render  # noqa: E402

OUT = ROOT / "out" / "features"
OUT.mkdir(parents=True, exist_ok=True)
ENV = {**os.environ, "PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"}
W = de.W


def tool(name, *args):
    r = subprocess.run([sys.executable, str(ROOT / "tools" / name), *map(str, args)], capture_output=True,
                       text=True, encoding="utf-8", errors="replace", cwd=ROOT, env=ENV)
    return r.returncode, r.stdout + r.stderr


def edit(src, name, ops, *extra):
    dst, plan = OUT / f"{name}.docx", OUT / f"{name}.json"
    dst.unlink(missing_ok=True)
    plan.write_text(json.dumps(ops, ensure_ascii=False), encoding="utf-8")
    rc, out = tool("docx_edit.py", src, dst, "--ops", plan, *extra)
    return rc, out, dst


def body(path):
    zin, trees = de.open_package(path)
    zin.close()
    return trees[de.MAIN]


def grid(tbl):
    return [int(g.get(de.q("w"))) for g in tbl.find(de.q("tblGrid"))]


def cells(tbl):
    return [[de.cell_text(tc) for tc in de.row_cells(tr)] for tr in de.table_rows(tbl)]


def picture(path, w=320, h=160, rgb=(0.2, 0.45, 0.7)):
    doc = fitz.open()
    page = doc.new_page(width=w, height=h)
    page.draw_rect(page.rect, color=None, fill=rgb)
    page.insert_text((20, h / 2), path.stem, fontsize=24, color=(1, 1, 1))
    page.get_pixmap(dpi=72).save(path)
    return path


class DocxFeatures(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        doc = new_document()
        doc.add_heading("Отчёт о продажах", 1)
        p = doc.add_paragraph(style="Body Text")
        p.add_run("Срок ")
        r = p.add_run("оплаты")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        r.font.name = "Georgia"
        p.add_run(" составляет 30 дней.")
        add_table(doc, [["Регион", "План", "Факт", "Δ"], ["Москва", "120,0", "131,5", "+9,6 %"],
                        ["Казань", "80,0", "77,2", "−3,5 %"], ["Итого", "200,0", "208,7", "+4,4 %"]],
                  widths=[4, 2, 2, 2], caption="Таблица 1. Продажи")
        doc.add_paragraph("Заключительный абзац.", style="Body Text")
        cls.src = OUT / "base.docx"
        doc.save(cls.src)

    def test_add_and_delete_column_keep_width_and_look(self):
        before = next(body(self.src).iter(de.q("tbl")))
        rc, out, dst = edit(self.src, "col_add", [{"op": "add_col", "table": 0, "after": 1,
                                                   "values": ["Прогноз", "125,0", "82,0", "207,0"]}])
        self.assertEqual(rc, 0, out)
        tbl = next(body(dst).iter(de.q("tbl")))
        self.assertEqual([r[2] for r in cells(tbl)], ["Прогноз", "125,0", "82,0", "207,0"])
        self.assertEqual(len(grid(tbl)), 5)
        self.assertEqual(sum(grid(tbl)), sum(grid(before)))  # the table did not get wider than the page
        for tr in de.table_rows(tbl):                         # cell widths follow the grid
            self.assertEqual([int(tc.find(f"{de.q('tcPr')}/{de.q('tcW')}").get(de.q("w"))) for tc in de.row_cells(tr)], grid(tbl))
        head = de.row_cells(de.table_rows(tbl)[0])
        self.assertIsNotNone(head[2].find(f".//{de.q('rPr')}/{de.q('b')}"))            # header look cloned
        self.assertIsNotNone(head[2].find(f"{de.q('tcPr')}/{de.q('tcBorders')}"))
        rc, out, back = edit(dst, "col_del", [{"op": "del_col", "table": 0, "col": 2}])
        self.assertEqual(rc, 0, out)
        tbl2 = next(body(back).iter(de.q("tbl")))
        self.assertEqual(cells(tbl2), cells(before))
        self.assertEqual(sum(grid(tbl2)), sum(grid(before)))
        self.assertTrue(all(abs(a - b) <= 2 for a, b in zip(grid(tbl2), grid(before))))

    def test_column_through_merged_cell_grows_the_merge(self):
        rc, out, merged = edit(self.src, "merge", [{"op": "merge_cells", "table": 0, "row": 3, "col": 0, "cols": 2}])
        self.assertEqual(rc, 0, out)
        tbl = next(body(merged).iter(de.q("tbl")))
        last = de.row_cells(de.table_rows(tbl)[3])
        self.assertEqual(len(last), 3)
        self.assertEqual(last[0].find(f"{de.q('tcPr')}/{de.q('gridSpan')}").get(de.q("val")), "2")
        self.assertIn("Итого", de.cell_text(last[0]))
        self.assertIn("200,0", de.cell_text(last[0]))  # text of the absorbed cell is kept
        rc, out, dst = edit(merged, "merge_col", [{"op": "add_col", "table": 0, "after": 0, "values": ["Код", "77", "16", ""]}])
        self.assertEqual(rc, 0, out)
        tbl = next(body(dst).iter(de.q("tbl")))
        self.assertEqual(de.row_cells(de.table_rows(tbl)[3])[0].find(f"{de.q('tcPr')}/{de.q('gridSpan')}").get(de.q("val")), "3")
        rc, out, _ = edit(merged, "merge_bad", [{"op": "add_col", "table": 0, "after": 0, "values": ["Код", "77", "16", "x"]}])
        self.assertEqual(rc, 1)
        self.assertIn("merged cell", out)

    def test_vertical_merge_and_widths(self):
        rc, out, dst = edit(self.src, "vmerge", [{"op": "merge_cells", "table": 0, "row": 1, "col": 0, "rows": 2},
                                                 {"op": "col_widths", "table": 0, "widths": [1, 1, 1, 1]}])
        self.assertEqual(rc, 0, out)
        tbl = next(body(dst).iter(de.q("tbl")))
        rows = de.table_rows(tbl)
        marks = [de.row_cells(rows[i])[0].find(f"{de.q('tcPr')}/{de.q('vMerge')}") for i in (1, 2)]
        self.assertEqual(marks[0].get(de.q("val")), "restart")
        self.assertIsNotNone(marks[1])
        self.assertIn("Казань", de.cell_text(de.row_cells(rows[1])[0]))
        g = grid(tbl)
        self.assertLessEqual(max(g) - min(g), 2)
        Document(str(dst))  # opens

    def test_format_changes_only_what_is_named(self):
        rc, out, dst = edit(self.src, "format", [
            {"op": "format", "find": "30 дней", "bold": True, "highlight": "yellow"},
            {"op": "format", "find": "оплаты", "italic": True},
            {"op": "format", "table": 0, "row": 0, "fill": "FFF2CC", "color": "1F4E79", "align": "center"},
            {"op": "format", "table": 0, "col": -1, "size": 9},
            {"op": "format", "block": "Заключительный абзац", "align": "right", "space_before": 12, "font": "Georgia"}])
        self.assertEqual(rc, 0, out)
        doc = Document(str(dst))
        para = next(p for p in doc.paragraphs if p.text.startswith("Срок"))
        runs = {r.text: r for r in para.runs}
        self.assertEqual(para.text, "Срок оплаты составляет 30 дней.")
        self.assertTrue(runs["30 дней"].bold)
        self.assertFalse(runs[" составляет "].bold)
        word = runs["оплаты"]  # the formats it already had are untouched
        self.assertTrue(word.bold and word.italic)
        self.assertEqual(str(word.font.color.rgb), "C00000")
        self.assertEqual(word.font.name, "Georgia")
        table = doc.tables[0]
        for c in table.rows[0].cells:
            self.assertIn('w:fill="FFF2CC"', c._tc.xml)
            self.assertEqual(str(c.paragraphs[0].runs[0].font.color.rgb), "1F4E79")
            self.assertTrue(c.paragraphs[0].runs[0].bold)  # header bold survived the recolouring
        self.assertEqual(table.rows[2].cells[3].paragraphs[0].runs[0].font.size, Pt(9))
        self.assertIsNone(table.rows[2].cells[0].paragraphs[0].runs[0].font.size)
        last = next(p for p in doc.paragraphs if p.text.startswith("Заключительный"))
        self.assertEqual(last.paragraph_format.space_before, Pt(12))
        self.assertEqual(last.runs[0].font.name, "Georgia")

    def test_tracked_format_records_old_properties(self):
        rc, out, dst = edit(self.src, "format_tracked", [{"op": "format", "find": "оплаты", "underline": True},
                                                         {"op": "format", "table": 0, "row": 1, "fill": "E2EFDA"}],
                            "--track", "Reviewer")
        self.assertEqual(rc, 0, out)
        root = body(dst)
        change = next(root.iter(de.q("rPrChange")))
        self.assertEqual(change.get(de.q("author")), "Reviewer")
        self.assertIsNotNone(change.find(f"{de.q('rPr')}/{de.q('b')}"))      # what it looked like before
        self.assertIsNone(change.find(f"{de.q('rPr')}/{de.q('u')}"))
        self.assertEqual(len(list(root.iter(de.q("tcPrChange")))), 4)
        Document(str(dst))

    def test_bad_format_is_refused(self):
        for op, why in (({"op": "format", "find": "оплаты"}, "at least one"),
                        ({"op": "format", "find": "оплаты", "color": "red"}, "hex"),
                        ({"op": "format", "find": "оплаты", "fill": "FFFFFF"}, "paragraphs/cells"),
                        ({"op": "format", "block": "#1", "style": "No Such Style"}, "not in document"),
                        ({"op": "format", "find": "нет такого", "bold": True}, "0 matches")):
            rc, out, dst = edit(self.src, "format_bad", [op])
            self.assertEqual(rc, 1, out)
            self.assertIn(why, out)
            self.assertFalse(dst.exists())

    def test_pictures_insert_replace_and_stay_in_the_column(self):
        chart = picture(OUT / "chart.png", 640, 320)
        picture(OUT / "logo.png", 200, 200, (0.7, 0.2, 0.2))
        rc, out, dst = edit(self.src, "image", [{"op": "insert_image", "after": "Таблица 1. Продажи", "file": "chart.png", "width_cm": 10}])
        self.assertEqual(rc, 0, out)
        doc = Document(str(dst))
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertEqual(round(doc.inline_shapes[0].width.cm, 1), 10.0)
        self.assertEqual(round(doc.inline_shapes[0].height.cm, 1), 5.0)   # aspect kept
        with zipfile.ZipFile(dst) as z:
            media = [n for n in z.namelist() if n.startswith("word/media/")]
            self.assertEqual(len(media), 1)
            self.assertEqual(z.read(media[0]), chart.read_bytes())
            self.assertIn('Extension="png"', z.read("[Content_Types].xml").decode())
        rc, out, swapped = edit(dst, "image_swap", [{"op": "replace_image", "image": 0, "file": "logo.png"}])
        self.assertEqual(rc, 0, out)
        shape = Document(str(swapped)).inline_shapes[0]
        self.assertLessEqual(shape.width.cm, 10.01)                       # fitted into the old box
        self.assertEqual(round(shape.width.cm, 1), round(shape.height.cm, 1))
        rc, out, _ = edit(self.src, "image_wide", [{"op": "insert_image", "after": "#0", "file": "chart.png", "width_cm": 40}])
        self.assertEqual(rc, 1)
        self.assertIn("wider than the text area", out)
        rc, out = tool("qa.py", swapped)
        self.assertEqual(rc, 0, out[-600:])

    def test_paragraphs_inside_cells_are_addressable(self):
        cell = {"table": 0, "row": 1, "col": 0}
        rc, out, dst = edit(self.src, "cell_para", [{"op": "insert", "after": cell, "text": "в т. ч. область"},
                                                    {"op": "format", "block": {**cell, "para": 1}, "italic": True}])
        self.assertEqual(rc, 0, out)
        c = Document(str(dst)).tables[0].rows[1].cells[0]
        self.assertEqual([p.text for p in c.paragraphs], ["Москва", "в т. ч. область"])
        self.assertTrue(c.paragraphs[1].runs[0].italic)
        rc, out, back = edit(dst, "cell_para_del", [{"op": "delete", "block": {**cell, "para": 1}}])
        self.assertEqual(rc, 0, out)
        self.assertEqual([p.text for p in Document(str(back)).tables[0].rows[1].cells[0].paragraphs], ["Москва"])
        rc, out, _ = edit(self.src, "cell_para_last", [{"op": "delete", "block": cell}])
        self.assertEqual(rc, 1)
        self.assertIn("final paragraph", out)

    def test_content_controls_are_looked_through(self):
        """Templates wrap blocks, rows and cells in w:sdt; they must stay editable."""
        doc = Document(str(self.src))
        b = doc.element.body
        para = next(p for p in b.iter(de.q("p")) if "Заключительный" in "".join(p.itertext()))
        sdt = etree.fromstring(f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Block"/></w:sdtPr><w:sdtContent/></w:sdt>')
        para.addprevious(sdt)
        sdt[1].append(para)
        row = list(b.iter(de.q("tr")))[2]
        rsdt = etree.fromstring(f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Row"/></w:sdtPr><w:sdtContent/></w:sdt>')
        row.addprevious(rsdt)
        rsdt[1].append(row)
        first = next(list(b.iter(de.q("tr")))[1].iter(de.q("tc")))   # a cell wrapped on its own, like a form field
        csdt = etree.fromstring(f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Cell"/></w:sdtPr><w:sdtContent/></w:sdt>')
        first.addprevious(csdt)
        csdt[1].append(first)
        src = OUT / "sdt.docx"
        doc.save(src)
        rc, out = tool("docx_edit.py", src, "--tables")
        self.assertIn("r2: Казань", out)
        rc, out, dst = edit(src, "sdt_edit", [{"op": "set_text", "block": "Заключительный абзац", "text": "Новый финал."},
                                              {"op": "cell", "table": 0, "row": 2, "col": 1, "text": "81,0"},
                                              {"op": "add_row", "table": 0, "after": 2, "values": ["Самара", "60,0", "64,9", "+8,2 %"]},
                                              {"op": "del_row", "table": 0, "row": 2}])
        self.assertEqual(rc, 0, out)
        root = body(dst)
        self.assertEqual([r[0] for r in cells(next(root.iter(de.q("tbl"))))], ["Регион", "Москва", "Самара", "Итого"])
        self.assertIn("Новый финал.", "".join(root.itertext()))
        self.assertEqual(len(list(root.iter(de.q("sdt")))), 2)  # the emptied row control went with its row
        rc, out, more = edit(dst, "sdt_cellrow", [{"op": "add_row", "table": 0, "after": 1, "values": ["Тула", "1", "2", "3"]},
                                                  {"op": "add_col", "table": 0, "after": 0, "values": ["Код", "77", "71", "63", ""]}])
        self.assertEqual(rc, 0, out)
        got = cells(next(body(more).iter(de.q("tbl"))))
        self.assertEqual(got[2], ["Тула", "71", "1", "2", "3"])       # cloned from a row whose first cell is wrapped
        self.assertEqual(got[1][:2], ["Москва", "77"])
        self.assertEqual(len(list(body(more).iter(de.q("sdt")))), 2)  # the clone did not duplicate the control
        Document(str(dst))

    def test_document_part_with_another_name(self):
        """Some producers store the body in word/document2.xml; the package says where."""
        src = OUT / "altmain.docx"
        with zipfile.ZipFile(self.src) as zin, zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data, name = zin.read(info.filename), info.filename
                if name == "word/document.xml":
                    name = "word/document2.xml"
                elif name == "word/_rels/document.xml.rels":
                    name = "word/_rels/document2.xml.rels"
                elif name in ("_rels/.rels", "[Content_Types].xml"):
                    data = data.replace(b"word/document.xml", b"word/document2.xml")
                zout.writestr(name, data)
        picture(OUT / "chart.png", 640, 320)
        rc, out, dst = edit(src, "altmain_edit", [{"op": "replace", "find": "30 дней", "replace": "45 дней"},
                                                  {"op": "insert_image", "after": "Заключительный абзац", "file": "chart.png", "width_cm": 8}])
        self.assertEqual(rc, 0, out)
        with zipfile.ZipFile(dst) as z:
            self.assertIn("word/document2.xml", z.namelist())
            self.assertNotIn("word/document.xml", z.namelist())
            self.assertTrue("media/harness1.png" in z.read("word/_rels/document2.xml.rels").decode("utf8"))
        doc = Document(str(dst))
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertTrue(any(p.text == "Срок оплаты составляет 45 дней." for p in doc.paragraphs))

    def test_listing_and_wrong_inputs(self):
        rc, out = tool("docx_edit.py", self.src, "--blocks", "--full")
        self.assertEqual(rc, 0)
        self.assertIn("r3: Итого | 200,0 | 208,7 | +4,4 %", out)
        self.assertIn("Срок оплаты составляет 30 дней.", out)
        rc, out = tool("docx_edit.py", ROOT / "README.md", "--blocks")
        self.assertEqual(rc, 1)
        self.assertIn("convert first", out)
        bomb = OUT / "xxe.docx"
        secret = OUT / "secret.txt"
        secret.write_text("TOP-SECRET-CONTENT", encoding="utf-8")
        with zipfile.ZipFile(self.src) as zin, zipfile.ZipFile(bomb, "w") as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == "word/document.xml":
                    text = data.decode("utf8")
                    head, rest = text.split("?>", 1)
                    data = (head + f'?><!DOCTYPE d [<!ENTITY x SYSTEM "{secret.as_uri()}">]>'
                            + rest.replace("Заключительный абзац.", "Заключительный &x; абзац.")).encode("utf8")
                zout.writestr(info, data)
        rc, out = tool("docx_edit.py", bomb, "--blocks", "--full")
        self.assertNotIn("TOP-SECRET-CONTENT", out)  # external entities are never pulled in


TYP = '''#import "/templates/report.typ": *
#show: report.with(title: "Отчёт", author: "Тест", date: "2 октября 2026")
= Продажи
Первый абзац отчёта: он достаточно длинный, чтобы занять две строки набора, и служит якорем для
вставки и удаления блоков в PDF без исходника.

#figure(table(columns: (1fr, auto, auto), align: (left, right, right),
  table.header[Регион][План][Факт],
  [Москва], [120,0], [131,5],
  [Казань], [80,0], [77,2],
  [Самара], [60,0], [64,9],
), caption: [Продажи по регионам])

Второй абзац после таблицы.
'''


def lines(doc, pno=0):
    return {pf.norm_ws(" ".join(c[2] for c in r.cells)): r for r in pf.visual_rows(doc[pno])}


class PdfFlow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        typ = OUT / "flow.typ"
        typ.write_text(TYP, encoding="utf-8")
        cls.src, _ = render(typ, OUT / "flow_render")
        cls.base = fitz.open(cls.src)
        cls.rows = lines(cls.base)

    @classmethod
    def tearDownClass(cls):
        cls.base.close()

    def run_ops(self, name, ops):
        dst, plan = OUT / f"{name}.pdf", OUT / f"{name}.json"
        dst.unlink(missing_ok=True)
        plan.write_text(json.dumps(ops, ensure_ascii=False), encoding="utf-8")
        rc, out = tool("pdf_flow.py", self.src, dst, "--ops", plan)
        return rc, out, dst

    def test_add_row_looks_like_its_neighbour(self):
        rc, out, dst = self.run_ops("add_row", [{"op": "add_row", "page": 1, "after": "Казань", "values": ["Уфа", "40,0", "42,1"]}])
        self.assertEqual(rc, 0, out)
        with fitz.open(dst) as doc:
            rows = lines(doc)
            new, old, below = rows["Уфа 40,0 42,1"], rows["Казань 80,0 77,2"], rows["Самара 60,0 64,9"]
            pitch = self.rows["Самара 60,0 64,9"].baseline - self.rows["Казань 80,0 77,2"].baseline
            self.assertAlmostEqual(new.baseline - old.baseline, pitch, delta=0.6)
            self.assertAlmostEqual(below.baseline - self.rows["Самара 60,0 64,9"].baseline, pitch, delta=0.6)
            self.assertAlmostEqual(old.baseline, self.rows["Казань 80,0 77,2"].baseline, delta=0.01)  # above: untouched
            for a, b in zip(new.cells[1:], old.cells[1:]):
                self.assertAlmostEqual(a[1], b[1], delta=0.6)       # numbers stay right-aligned to the column
            a, b = new.cells[0][3][0], old.cells[0][3][0]
            self.assertEqual((a["font"].split("+")[-1], round(a["size"], 1), a["color"]),
                             (b["font"].split("+")[-1], round(b["size"], 1), b["color"]))
            footer = [r for t, r in rows.items() if t == "1 / 1"][0]
            self.assertAlmostEqual(footer.baseline, self.rows["1 / 1"].baseline, delta=0.01)       # footer stays
            self.assertTrue(all(f[1] != "n/a" for f in doc[0].get_fonts()))
        rc, out = tool("qa.py", dst, "--original", self.src, "--allow-reflow")
        self.assertEqual(rc, 0, out[-600:])

    def test_del_row_and_cell(self):
        rc, out, dst = self.run_ops("del_row", [{"op": "del_row", "page": 1, "row": "Казань"},
                                                {"op": "cell", "page": 1, "row": "Самара", "col": 2, "text": "164,9"}])
        self.assertEqual(rc, 0, out)
        with fitz.open(dst) as doc:
            rows = lines(doc)
            self.assertNotIn("Казань", doc[0].get_text())
            new = rows["Самара 60,0 164,9"]
            old = self.rows["Самара 60,0 64,9"]
            self.assertAlmostEqual(new.baseline, self.rows["Казань 80,0 77,2"].baseline, delta=0.6)  # moved up one row
            self.assertAlmostEqual(new.cells[2][1], old.cells[2][1], delta=0.6)                      # right edge kept
            self.assertLess(rows["Второй абзац после таблицы."].baseline, self.rows["Второй абзац после таблицы."].baseline)

    def test_insert_and_delete_block(self):
        text = "Вставленный абзац: набран шрифтом соседнего абзаца и сдвигает всё, что ниже."
        rc, out, dst = self.run_ops("insert", [{"op": "insert", "page": 1, "after": "Первый абзац отчёта", "text": text}])
        self.assertEqual(rc, 0, out)
        with fitz.open(dst) as doc:
            rows = lines(doc)
            new = next(r for t, r in rows.items() if t.startswith("Вставленный абзац"))
            anchor = next(r for t, r in rows.items() if t.startswith("Первый абзац"))
            self.assertTrue(pf.norm_ws(text)[-30:] in pf.norm_ws(doc[0].get_text()))
            self.assertLess(anchor.baseline, new.baseline)
            self.assertLess(new.baseline, rows["Регион План Факт"].baseline)
            like = anchor.spans[0]
            self.assertEqual((new.spans[0]["font"].split("+")[-1], round(new.spans[0]["size"], 1)),
                             (like["font"].split("+")[-1], round(like["size"], 1)))
            self.assertGreater(rows["Москва 120,0 131,5"].baseline, self.rows["Москва 120,0 131,5"].baseline + 10)
        rc, out, dst = self.run_ops("delete", [{"op": "delete", "page": 1, "block": "Первый абзац отчёта"}])
        self.assertEqual(rc, 0, out)
        with fitz.open(dst) as doc:
            self.assertNotIn("Первый абзац", doc[0].get_text())
            self.assertLess(lines(doc)["Москва 120,0 131,5"].baseline, self.rows["Москва 120,0 131,5"].baseline - 10)

    def test_insert_image_opens_room_for_it(self):
        picture(OUT / "chart.png", 640, 320)
        rc, out, dst = self.run_ops("image", [{"op": "insert_image", "page": 1, "after": "Второй абзац после таблицы", "file": "chart.png", "width_cm": 8}])
        self.assertEqual(rc, 0, out)
        with fitz.open(dst) as doc:
            info = doc[0].get_image_info()
            self.assertEqual(len(info), 1)
            box = fitz.Rect(info[0]["bbox"])
            self.assertAlmostEqual(box.width, 8 / 2.54 * 72, delta=1)
            self.assertAlmostEqual(box.height, box.width / 2, delta=1)
            anchor = lines(doc)["Второй абзац после таблицы."]
            self.assertGreater(box.y0, anchor.y1)
            self.assertAlmostEqual(anchor.baseline, self.rows["Второй абзац после таблицы."].baseline, delta=0.01)
        rc, out, dst = self.run_ops("image_wide", [{"op": "insert_image", "page": 1, "after": "Второй абзац после таблицы", "file": "chart.png", "width_cm": 40}])
        self.assertEqual(rc, 1)
        self.assertTrue("wider than the text column" in out)

    def test_format_colour_and_fill(self):
        rc, out, dst = self.run_ops("format", [{"op": "format", "page": 1, "row": "Казань", "fill": "FFF2CC"},
                                               {"op": "format", "page": 1, "row": "Самара", "col": 2, "color": "C00000"}])
        self.assertEqual(rc, 0, out)
        with fitz.open(dst) as doc:
            rows = lines(doc)
            cell = rows["Самара 60,0 64,9"].cells
            self.assertEqual({s["color"] for s in cell[2][3]}, {0xC00000})
            self.assertEqual({s["color"] for s in cell[1][3]}, {s["color"] for s in self.rows["Самара 60,0 64,9"].cells[1][3]})
            row = rows["Казань 80,0 77,2"]
            self.assertAlmostEqual(row.baseline, self.rows["Казань 80,0 77,2"].baseline, delta=0.01)   # nothing moved
            pix = doc[0].get_pixmap(clip=fitz.Rect(row.cells[0][1] + 20, row.y0, row.cells[1][0] - 20, row.y1), colorspace=fitz.csRGB)
            self.assertEqual(pix.pixel(pix.width // 2, pix.height // 2), (0xFF, 0xF2, 0xCC))        # filled behind the text
            self.assertEqual(doc[0].get_text("words"), doc[0].get_text("words"))
            self.assertEqual(sorted(w[4] for w in doc[0].get_text("words")), sorted(w[4] for w in self.base[0].get_text("words")))
        rc, out, _ = self.run_ops("format_bad", [{"op": "format", "page": 1, "row": "Казань", "bold": True}])
        self.assertEqual(rc, 1)
        self.assertTrue("need the source document" in out)

    def test_refusals_write_nothing(self):
        for name, ops, why in (
                ("r_ambiguous", [{"op": "del_row", "page": 1, "row": "0"}], "need exactly 1"),
                ("r_values", [{"op": "add_row", "page": 1, "after": "Казань", "values": ["Уфа"]}], "3 cells"),
                ("r_wide", [{"op": "cell", "page": 1, "row": "Казань", "col": 1, "text": "очень-очень длинное значение, которое не влезет в колонку"}], "does not fit"),
                ("r_room", [{"op": "open_gap", "page": 1, "y": round(self.rows["Второй абзац после таблицы."].y0 - 2, 1), "height": 700}], "no room"),
                ("r_slice", [{"op": "open_gap", "page": 1, "y": round(self.rows["Казань 80,0 77,2"].baseline - 3, 1), "height": 10}], "slice through text"),
                ("r_page", [{"op": "del_row", "page": 9, "row": "Казань"}], "page must be")):
            rc, out, dst = self.run_ops(name, ops)
            self.assertEqual(rc, 1, out)
            self.assertTrue(why in out, f"{name}: {out[-200:]}")
            self.assertFalse(dst.exists(), name)

    def test_relative_text_operators_move_as_a_unit(self):
        """Td / TD / T* / ' position lines relative to the previous one: moving a part of such a
        chain must not drag the lines that stay."""
        pdf = pikepdf.new()
        pdf.add_blank_page(page_size=(300, 400))
        page = pdf.pages[0]
        page.obj["/Resources"] = pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=pikepdf.Dictionary(
            Type=pikepdf.Name.Font, Subtype=pikepdf.Name.Type1, BaseFont=pikepdf.Name.Helvetica)))
        # baselines from the page top: alpha 50, beta 80, gamma 110, delta 150, epsilon 190
        stream = (b"q 0.5 w 40 285 m 260 285 l S 40 330 220 -100 re S Q "
                  b"BT /F1 12 Tf 30 TL 50 350 Td (alpha) Tj 0 -30 Td (beta) Tj T* (gamma) Tj "
                  b"0 -40 TD (delta) Tj (epsilon) ' ET")
        page.obj["/Contents"] = pdf.make_stream(stream)
        buf = io.BytesIO()
        pdf.save(buf)
        flow = pf.Flow(buf.getvalue())
        base = {w[4]: w[3] for w in flow.doc[0].get_text("words")}
        flow.open_gap(flow.doc[0], 90.0, 20.0)   # between "beta" and "gamma"
        after = {w[4]: w[3] for w in flow.doc[0].get_text("words")}
        for word, shift in (("alpha", 0), ("beta", 0), ("gamma", 20), ("delta", 20), ("epsilon", 20)):
            self.assertAlmostEqual(after[word] - base[word], shift, delta=0.05, msg=word)
        box = [d["rect"] for d in flow.doc[0].get_drawings() if d["rect"].height > 50][0]
        self.assertAlmostEqual(box.height, 120, delta=0.6)   # the frame crossing the cut grew by the gap
        flow.close_band(flow.doc[0], 60.0, 90.0)   # removes "beta", closes the 30pt
        final = {w[4]: w[3] for w in flow.doc[0].get_text("words")}
        self.assertNotIn("beta", final)
        self.assertAlmostEqual(final["gamma"] - base["gamma"], -10, delta=0.05)
        self.assertAlmostEqual(final["alpha"], base["alpha"], delta=0.05)
        flow.doc.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
