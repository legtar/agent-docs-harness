"""House style for NEW .docx files.

    python tools/docx_kit.py reference              # (re)build templates/reference.docx from TOKENS
    python tools/docx_kit.py md IN.md OUT.docx [--toc]   # Markdown -> styled docx (pandoc + table fix-up)

From Python (python-docx authoring):
    from docx_kit import new_document, add_table, add_toc
    doc = new_document()                      # styles from reference.docx, empty body
    doc.add_heading("Заголовок", 1); doc.add_paragraph("Текст", style="Body Text")
    add_table(doc, [["Показатель", "2025"], ["Выручка", "412,0"]], widths=[12, 4], caption="Таблица 1. …")
    doc.save("out/x.docx")                     # then: python tools/qa.py out/x.docx

Layout-stability rules baked in: fonts every Office install has (no substitution at the
recipient), headings keep-with-next, widow control, fixed-width tables that never exceed the
text column, header rows repeat, rows never split across pages, no hyphenation (renderer-
dependent), body left-aligned (justification reflows differently in Word vs LibreOffice).
"""
import pathlib
import re
import shutil
import subprocess
import sys
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
REFERENCE = ROOT / "templates" / "reference.docx"

TOKENS = {
    "body_font": "Calibri", "head_font": "Georgia", "lang": "ru-RU",
    "body_size": 11, "h1": 18, "h2": 14, "h3": 12, "title": 26,
    "ink": "1A1A1A", "accent": "1F4E79", "muted": "6B7280", "rule": "BFC5CD", "tint": "F3F6FA",
    "margins_cm": (2.2, 2.2, 2.0, 2.4),  # left, right, top, bottom
    "page": (21.0, 29.7),                # A4
}
NUM = re.compile(r"^[\s+\-−–]?[\d\s\u00a0.,]+\s*(%|‰|₽|\$|€|п\.п\.|млн|млрд|тыс\.?)?$")


def pandoc() -> str:
    for c in (shutil.which("pandoc"), pathlib.Path.home() / "AppData/Local/Pandoc/pandoc.exe",
              pathlib.Path("C:/Program Files/Pandoc/pandoc.exe")):
        if c and pathlib.Path(c).exists():
            return str(c)
    sys.exit("pandoc not found: winget install -e --id JohnMacFarlane.Pandoc")


# ---------- low-level OOXML helpers ----------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


# Schema child order (wml.xsd). Out-of-order children = Word "unreadable content" error.
ORDER = {
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc "
             "tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "settings": "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
                "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData "
                "embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins alignBordersAndEdges "
                "bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors "
                "hideGrammaticalErrors activeWritingStyle proofState formsDesign attachedTemplate linkStyles "
                "stylePaneFormatFilter stylePaneSortMethod documentType mailMerge revisionView trackRevisions "
                "doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride styleLockTheme "
                "styleLockQFSet defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone "
                "doNotHyphenateCaps showEnvelope summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders "
                "bookFoldRevPrinting bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing "
                "drawingGridVerticalSpacing displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery "
                "doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin drawingGridVerticalOrigin "
                "doNotShadeFormData noPunctuationKerning characterSpacingControl printTwoOnOne strictFirstAndLastChars "
                "noLineBreaksAfter noLineBreaksBefore savePreviewPicture doNotValidateAgainstSchema saveInvalidXml "
                "ignoreMixedContent alwaysShowPlaceholderText doNotDemarcateInvalidXml saveXmlDataOnly "
                "useXSLTWhenSaving saveThroughXslt showXMLTags alwaysMergeEmptyNamespace updateFields "
                "hdrShapeDefaults footnotePr endnotePr compat docVars rsids",
}
ORDER = {k: v.split() for k, v in ORDER.items()}


def _local(e):
    return e.tag.rsplit("}", 1)[-1]


def _set(parent, tag, **attrs):
    """Replace-or-create a single child element, at its schema position."""
    for old in parent.findall(qn(tag)):
        parent.remove(old)
    e = _el(tag, **attrs)
    order, name = ORDER.get(_local(parent)), tag.split(":")[1]
    if order and name in order:
        for sib in parent:
            n = _local(sib)
            if n in order and order.index(n) > order.index(name):
                sib.addprevious(e)
                return e
    parent.append(e)
    return e


def _fonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rpr.insert(0, rf)
    for a in list(rf.attrib):
        if a.endswith("Theme"):  # theme fonts would override the explicit name
            del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def _style(doc, name, font=None, size=None, color=None, bold=None, before=None, after=None,
           keep_next=None, line=None, align=None):
    s = next((x for x in doc.styles if x.name and x.name.lower() == name.lower()), None)
    if s is None:
        return None
    if font:
        _fonts(s.element.get_or_add_rPr(), font)
    f = s.font
    if size: f.size = Pt(size)
    if color: f.color.rgb = RGBColor.from_string(color)
    if bold is not None: f.bold = bold
    pf = getattr(s, "paragraph_format", None)
    if pf is not None:
        if before is not None: pf.space_before = Pt(before)
        if after is not None: pf.space_after = Pt(after)
        if keep_next is not None:
            pf.keep_with_next = keep_next
            pf.keep_together = keep_next
        if line is not None: pf.line_spacing = line
        if align is not None: pf.alignment = align
        pf.widow_control = True
    return s


def _table_style(doc, T):
    """Booktabs look for pandoc's 'Table' style: rules top/bottom/under header, hairlines inside."""
    try:
        st = doc.styles["Table"].element
    except KeyError:
        return
    tblPr = st.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = _el("w:tblPr"); st.append(tblPr)
    b = _set(tblPr, "w:tblBorders")
    for side, sz, col in (("top", 8, T["ink"]), ("bottom", 8, T["ink"]), ("insideH", 4, T["rule"])):
        b.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": sz, "w:space": 0, "w:color": col}))
    for side in ("left", "right", "insideV"):
        b.append(_el(f"w:{side}", **{"w:val": "nil"}))
    m = _set(tblPr, "w:tblCellMar")
    for side, v in (("top", 60), ("left", 100), ("bottom", 60), ("right", 100)):
        m.append(_el(f"w:{side}", **{"w:w": v, "w:type": "dxa"}))
    for old in st.findall(qn("w:tblStylePr")):
        st.remove(old)
    fr = _el("w:tblStylePr", **{"w:type": "firstRow"})
    rpr = _el("w:rPr"); rpr.append(_el("w:b")); rpr.append(_el("w:sz", **{"w:val": 19}))
    fr.append(rpr)
    tcpr = _el("w:tcPr"); tb = _el("w:tcBorders")
    tb.append(_el("w:bottom", **{"w:val": "single", "w:sz": 6, "w:space": 0, "w:color": T["ink"]}))
    tcpr.append(tb)
    fr.append(tcpr)
    st.append(fr)


def build_reference(path=REFERENCE, T=TOKENS):
    path.parent.mkdir(parents=True, exist_ok=True)
    raw = subprocess.run([pandoc(), "--print-default-data-file", "reference.docx"],
                         capture_output=True, check=True).stdout
    path.write_bytes(raw)
    doc = Document(str(path))
    # docDefaults: explicit fonts + language (spell-check / quotes), no theme indirection.
    rpr_def = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _fonts(rpr_def, T["body_font"])
    _set(rpr_def, "w:lang", **{"w:val": T["lang"], "w:eastAsia": T["lang"], "w:bidi": "ar-SA"})
    _style(doc, "Normal", T["body_font"], T["body_size"], T["ink"], after=6, line=1.2)
    for n in ("Body Text", "First Paragraph", "Compact"):
        _style(doc, n, before=0, after=6 if n != "Compact" else 2, align=WD_ALIGN_PARAGRAPH.LEFT)
    for lvl, size, before, after in ((1, T["h1"], 20, 8), (2, T["h2"], 14, 6), (3, T["h3"], 10, 4)):
        _style(doc, f"Heading {lvl}", T["head_font"], size, T["accent"] if lvl < 3 else T["ink"],
               True, before, after, keep_next=True, line=1.1)
    for n in ("Heading 4", "Heading 5", "Heading 6"):
        _style(doc, n, T["body_font"], T["body_size"], T["ink"], True, 8, 2, keep_next=True)
    left = WD_ALIGN_PARAGRAPH.LEFT
    _style(doc, "Title", T["head_font"], T["title"], T["accent"], True, 0, 6, keep_next=True, line=1.05, align=left)
    _style(doc, "Subtitle", T["body_font"], 13, T["muted"], False, 0, 14, keep_next=True, align=left)
    for n in ("Author", "Date", "Abstract"):
        _style(doc, n, T["body_font"], 10, T["muted"], after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    for n in ("Caption", "Table Caption", "Image Caption"):
        _style(doc, n, T["body_font"], 9, T["muted"], False, 4, 6, keep_next=(n == "Table Caption"))
    _style(doc, "Block Text", size=T["body_size"], color=T["muted"])
    _table_style(doc, T)
    # Page: size, margins, page number "N / M" in the footer.
    for s in doc.sections:
        s.page_width, s.page_height = Cm(T["page"][0]), Cm(T["page"][1])
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = (Cm(v) for v in T["margins_cm"])
        p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for part in ("PAGE", " / ", "NUMPAGES"):
            r = p.add_run(part if part == " / " else "1")
            r.font.size, r.font.color.rgb = Pt(8), RGBColor.from_string(T["muted"])
            if part != " / ":  # wrap the run in a live field
                fld = _el("w:fldSimple", **{"w:instr": f" {part} "})
                r._r.addprevious(fld)
                fld.append(r._r)
    doc.save(str(path))
    _theme_fonts(path, T)
    return path


def _theme_fonts(path, T):
    """Theme major/minor fonts -> our fonts, so nothing in the file can fall back to Aptos etc."""
    src = zipfile.ZipFile(path)
    items = [(i, src.read(i.filename)) for i in src.infolist()]
    src.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename.startswith("word/theme/"):
                x = data.decode("utf8")
                x = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', lambda m: m.group(1) + T["head_font"], x)
                x = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', lambda m: m.group(1) + T["body_font"], x)
                data = x.encode("utf8")
            z.writestr(info, data)


# ---------- authoring helpers ----------
def new_document(reference=REFERENCE):
    if not pathlib.Path(reference).exists():
        build_reference(pathlib.Path(reference))
    doc = Document(str(reference))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return doc


def text_width_dxa(doc):
    s = doc.sections[-1]
    return int((s.page_width - s.left_margin - s.right_margin) / 635)  # EMU -> dxa


def is_num(s: str) -> bool:
    return bool(s.strip()) and bool(NUM.match(s.strip()))


def polish_table(tbl, width_dxa, widths=None, header=True):
    """Fixed layout at full column width, repeat header, rows don't split, numbers right."""
    t = tbl._tbl
    tblPr = t.tblPr
    _set(tblPr, "w:tblW", **{"w:w": width_dxa, "w:type": "dxa"})
    _set(tblPr, "w:tblLayout", **{"w:type": "fixed"})
    ncols = len(t.tblGrid.findall(qn("w:gridCol")))
    if widths:
        tot = sum(widths)
        cols = [int(width_dxa * w / tot) for w in widths]
    else:
        # proportional to the longest text in each column, clamped so no column starves
        lens = [max((len(c.text) for c in col.cells), default=1) for col in tbl.columns]
        lens = [min(max(n, 4), 40) for n in lens]
        cols = [int(width_dxa * n / sum(lens)) for n in lens]
    cols[-1] += width_dxa - sum(cols)
    for g, w in zip(t.tblGrid.findall(qn("w:gridCol")), cols):
        g.set(qn("w:w"), str(w))
    # Direct borders, not table-style conditionals: identical in Word, LibreOffice, Google Docs.
    T = TOKENS
    b = _set(tblPr, "w:tblBorders")
    for side, sz, col in (("top", 8, T["ink"]), ("bottom", 8, T["ink"]), ("insideH", 4, T["rule"])):
        b.append(_el(f"w:{side}", **{"w:val": "single", "w:sz": sz, "w:space": 0, "w:color": col}))
    for side in ("left", "right", "insideV"):
        b.append(_el(f"w:{side}", **{"w:val": "nil"}))
    m = _set(tblPr, "w:tblCellMar")
    for side, v in (("top", 50), ("left", 100), ("bottom", 50), ("right", 100)):
        m.append(_el(f"w:{side}", **{"w:w": v, "w:type": "dxa"}))
    rows = list(tbl.rows)
    body = rows[1:] if header else rows
    numeric = [bool(body) and all(ci < len(r.cells) and (is_num(r.cells[ci].text) or not r.cells[ci].text.strip())
                                  for r in body) and any(ci < len(r.cells) and is_num(r.cells[ci].text) for r in body)
               for ci in range(ncols)]
    for ri, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        _set(trPr, "w:cantSplit")
        if header and ri == 0:
            _set(trPr, "w:tblHeader")
        for ci, cell in enumerate(row.cells[:ncols]):
            tcPr = cell._tc.get_or_add_tcPr()
            _set(tcPr, "w:tcW", **{"w:w": cols[min(ci, ncols - 1)], "w:type": "dxa"})
            if header and ri == 0:
                bd = _set(tcPr, "w:tcBorders")
                bd.append(_el("w:bottom", **{"w:val": "single", "w:sz": 6, "w:space": 0, "w:color": T["ink"]}))
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                # No orphan rows: short tables stay whole; long ones keep head+3 rows and the last
                # 3 rows together, so a page break can only fall in the middle.
                n = len(rows)
                p.paragraph_format.keep_with_next = ri < n - 1 and (n <= 8 or ri <= 3 or ri >= n - 3)
                if numeric[ci]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if header and ri == 0:
                    for r in p.runs:
                        r.bold = True
                        r.font.size = Pt(T["body_size"] - 1.5)
    return tbl


def add_table(doc, rows, widths=None, caption=None, header=True):
    """rows: list of lists of str. widths: relative column weights."""
    if caption:
        doc.add_paragraph(caption, style="Table Caption")
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = doc.styles["Table"]
    for r, data in zip(tbl.rows, rows):
        for c, v in zip(r.cells, data):
            c.text = str(v)
    return polish_table(tbl, text_width_dxa(doc), widths, header)


def add_toc(doc, title="Содержание", levels="1-3"):
    doc.add_paragraph(title, style="Heading 1").paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    r = p.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    it = _el("w:instrText"); it.text = f' TOC \\o "{levels}" \\h \\z \\u '
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r._r.append(it)
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    t = _el("w:t"); t.text = "Обновите поле (F9)"; r._r.append(t)
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    # Word refreshes fields on open; LibreOffice render fills the TOC itself.
    _set(doc.settings.element, "w:updateFields", **{"w:val": "true"})


def md_to_docx(md, out, toc=False):
    args = [pandoc(), str(md), "-o", str(out), "--reference-doc", str(REFERENCE if REFERENCE.exists() else build_reference())]
    if toc:
        args += ["--toc", "--toc-depth=3", "-M", "toc-title=Содержание"]
    subprocess.run(args, check=True)
    doc = Document(str(out))
    w = text_width_dxa(doc)
    for tbl in doc.tables:
        polish_table(tbl, w)
    doc.save(str(out))
    return out


if __name__ == "__main__":
    cmd = sys.argv[1:2]
    if cmd == ["reference"]:
        print(build_reference())
    elif cmd == ["md"] and len(sys.argv) >= 4:
        print(md_to_docx(sys.argv[2], sys.argv[3], "--toc" in sys.argv))
    else:
        sys.exit(__doc__)
